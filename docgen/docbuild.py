# -*- coding: utf-8 -*-
"""
Motor de construccion del Documento Final MedResQ en formato .docx

Reproduce el formato solicitado para reportes de residencia profesional:
  - Times New Roman 12, interlineado 1.5, texto justificado
  - Sangria de primera linea de 1.25 cm
  - Margenes de 2.5 cm
  - Estilos Titulo 1 (16 negrita), Titulo 2 (14 negrita), Titulo 3 (12 negrita)
  - Figura 10 pt centrado
  - Indice automatico (campo TOC), indice de figuras e indice de tablas
  - Encabezado con el nombre del reporte y numeracion de pagina abajo a la derecha
  - Saltos de seccion con numeracion romana en preliminares y arabiga en el cuerpo
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(BASE, "img")

NAVY = "0B0E2D"
NAVY_SOFT = "E8EFF9"
GREY_SOFT = "F4F6F9"
LINE = "C9D0DD"

TITULO_REPORTE = "Documento Final del Proyecto MedResQ"
ALUMNAS = "Rivera Carrillo Ingrid Guadalupe   •   Carrete Ruiz Abril Carolina"
INSTITUCION = "Universidad Politecnica de Durango"


# ----------------------------------------------------------------------- utilidades XML
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def _fuente(rpr_owner, nombre="Times New Roman"):
    rpr = rpr_owner.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), nombre)


def _sombra(celda, hexcolor):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _layout_fijo(tabla):
    tblPr = tabla._tbl.tblPr
    for viejo in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(viejo)
    tblPr.append(_el("w:tblLayout", type="fixed"))


def _bordes_tabla(tabla, color=LINE, sz=6):
    tblPr = tabla._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def _repetir_encabezado(fila):
    trPr = fila._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader", val="true"))


def _no_partir_fila(fila):
    trPr = fila._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _campo(parrafo, instruccion, texto_provisional="1"):
    """Inserta un campo de Word (PAGE, TOC, NUMPAGES...)."""
    r1 = parrafo.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)

    r2 = parrafo.add_run()
    it = OxmlElement("w:instrText")
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    it.text = instruccion
    r2._r.append(it)

    r3 = parrafo.add_run()
    fs = OxmlElement("w:fldChar")
    fs.set(qn("w:fldCharType"), "separate")
    r3._r.append(fs)

    r4 = parrafo.add_run(texto_provisional)

    r5 = parrafo.add_run()
    fe = OxmlElement("w:fldChar")
    fe.set(qn("w:fldCharType"), "end")
    r5._r.append(fe)
    return r4


def _numeracion(seccion, fmt="decimal", inicio=1):
    sectPr = seccion._sectPr
    for viejo in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(viejo)
    sectPr.append(_el("w:pgNumType", fmt=fmt, start=inicio))


def _borde_inferior(parrafo, color=LINE, sz=6):
    pPr = parrafo._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "3")
    b.set(qn("w:color"), color)
    pbdr.append(b)
    pPr.append(pbdr)


def _borde_superior(parrafo, color=LINE, sz=6):
    pPr = parrafo._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:top")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "3")
    b.set(qn("w:color"), color)
    pbdr.append(b)
    pPr.append(pbdr)


def _sombra_parrafo(parrafo, hexcolor):
    pPr = parrafo._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    pPr.append(shd)


def _mantener_con_siguiente(parrafo):
    parrafo.paragraph_format.keep_with_next = True


# ======================================================================= clase principal
class DocumentoFinal:

    def __init__(self):
        self.doc = Document()
        self.capitulo = 0
        self._primero_en_seccion = True
        self.fig_n = 0
        self.tab_n = 0
        self.indice_figuras = []
        self.indice_tablas = []
        self._configurar_estilos()
        self._configurar_seccion_inicial()

    # ------------------------------------------------------------------ estilos
    def _configurar_estilos(self):
        d = self.doc
        st = d.styles

        normal = st["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        normal.font.color.rgb = RGBColor(0, 0, 0)
        _fuente(normal.element)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)

        def crear(nombre, size, bold, align, space_before, space_after,
                  color="000000", italic=False, indent=0, spacing=1.5,
                  fuente="Times New Roman", base=None, mayus=False):
            if nombre in [s.name for s in st]:
                s = st[nombre]
            else:
                from docx.enum.style import WD_STYLE_TYPE
                s = st.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
                if base:
                    s.base_style = st[base]
            s.font.name = fuente
            s.font.size = Pt(size)
            s.font.bold = bold
            s.font.italic = italic
            s.font.color.rgb = RGBColor.from_string(color)
            _fuente(s.element, fuente)
            if mayus:
                rpr = s.element.get_or_add_rPr()
                rpr.append(_el("w:caps", val="true"))
            p = s.paragraph_format
            p.alignment = align
            p.space_before = Pt(space_before)
            p.space_after = Pt(space_after)
            p.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.line_spacing = spacing
            p.first_line_indent = Cm(indent)
            p.keep_with_next = size >= 13
            return s

        # Titulos jerarquicos (los usa el indice automatico)
        crear("Heading 1", 16, True, WD_ALIGN_PARAGRAPH.LEFT, 0, 18, NAVY, mayus=True)
        crear("Heading 2", 14, True, WD_ALIGN_PARAGRAPH.LEFT, 16, 10, NAVY)
        crear("Heading 3", 12, True, WD_ALIGN_PARAGRAPH.LEFT, 12, 8, NAVY)
        crear("Heading 4", 12, True, WD_ALIGN_PARAGRAPH.LEFT, 10, 6, "1F2452", italic=True)

        # Cuerpo
        crear("Cuerpo", 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, indent=1.25)
        crear("CuerpoSinSangria", 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, indent=0)
        crear("Vineta", 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 3, indent=0)
        crear("Figura", 10, True, WD_ALIGN_PARAGRAPH.CENTER, 6, 0, "000000", spacing=1.0)
        crear("FiguraFuente", 10, False, WD_ALIGN_PARAGRAPH.CENTER, 0, 12, "3A3A3A",
              italic=True, spacing=1.0)
        crear("Tabla", 10, True, WD_ALIGN_PARAGRAPH.CENTER, 10, 4, "000000", spacing=1.0)
        crear("TablaTexto", 10, False, WD_ALIGN_PARAGRAPH.LEFT, 2, 2, "000000", spacing=1.0)
        crear("TablaEncabezado", 10, True, WD_ALIGN_PARAGRAPH.CENTER, 2, 2, "FFFFFF", spacing=1.0)
        crear("Codigo", 9, False, WD_ALIGN_PARAGRAPH.LEFT, 0, 0, "16204A",
              spacing=1.0, fuente="Consolas")
        crear("Portada", 18, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 12, NAVY, spacing=1.5, mayus=True)
        crear("PortadaTexto", 14, False, WD_ALIGN_PARAGRAPH.CENTER, 0, 8, "000000", spacing=1.5)
        crear("PortadaChico", 12, False, WD_ALIGN_PARAGRAPH.CENTER, 0, 6, "000000", spacing=1.5)
        crear("TituloSuelto", 16, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 18, NAVY, mayus=True)
        crear("Nota", 11, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 6, 10, "3A3A3A",
              italic=True, spacing=1.15)

        # Vinetas y numeracion con estilos de lista nativos
        for ln in ("List Bullet", "List Number"):
            s = st[ln]
            s.font.name = "Times New Roman"
            s.font.size = Pt(12)
            _fuente(s.element)
            p = s.paragraph_format
            p.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.line_spacing = 1.5
            p.space_after = Pt(3)
            p.left_indent = Cm(1.25)
            p.first_line_indent = Cm(-0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # ------------------------------------------------------------------ secciones
    def _margenes(self, s):
        s.page_width = Cm(21.59)   # carta
        s.page_height = Cm(27.94)
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)
        s.header_distance = Cm(1.25)
        s.footer_distance = Cm(1.25)

    def _configurar_seccion_inicial(self):
        s = self.doc.sections[0]
        self._margenes(s)
        # Portada: sin encabezado ni pie
        s.header.is_linked_to_previous = False
        s.footer.is_linked_to_previous = False
        for p in list(s.header.paragraphs):
            p.text = ""
        for p in list(s.footer.paragraphs):
            p.text = ""

    def nueva_seccion(self, numeracion="decimal", inicio=1, con_encabezado=True):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._margenes(s)
        _numeracion(s, numeracion, inicio)
        s.header.is_linked_to_previous = False
        s.footer.is_linked_to_previous = False
        if con_encabezado:
            self._pintar_encabezado(s)
            self._pintar_pie(s)
        else:
            for p in list(s.header.paragraphs):
                p.text = ""
            for p in list(s.footer.paragraphs):
                p.text = ""
        self._primero_en_seccion = True
        return s

    def _pintar_encabezado(self, seccion):
        hdr = seccion.header
        p1 = hdr.paragraphs[0]
        p1.text = ""
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p1.add_run(TITULO_REPORTE.upper())
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        _fuente(r._r)
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.0

        p2 = hdr.add_paragraph()
        p2.paragraph_format.line_spacing = 1.0
        p2.paragraph_format.space_after = Pt(2)
        ancho = seccion.page_width - seccion.left_margin - seccion.right_margin
        p2.paragraph_format.tab_stops.add_tab_stop(Emu(int(ancho)), WD_TAB_ALIGNMENT.RIGHT)
        r2 = p2.add_run(ALUMNAS + "\t" + INSTITUCION)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor.from_string("3A3A3A")
        _borde_inferior(p2)

    def _pintar_pie(self, seccion):
        ftr = seccion.footer
        p = ftr.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.line_spacing = 1.0
        _borde_superior(p)
        r = p.add_run("")
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        run = _campo(p, "PAGE  \\* MERGEFORMAT")
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(NAVY)

    # ------------------------------------------------------------------ bloques de texto
    def h1(self, texto, numerar=True):
        """Titulo de capitulo. Comienza en pagina nueva."""
        if numerar:
            self.capitulo += 1
            self.fig_n = 0
            self.tab_n = 0
        if not self._primero_en_seccion:
            self.salto_pagina()
        self._primero_en_seccion = False
        p = self.doc.add_paragraph(texto, style="Heading 1")
        _borde_inferior(p, NAVY, 8)
        return p

    def h2(self, texto):
        return self.doc.add_paragraph(texto, style="Heading 2")

    def h3(self, texto):
        return self.doc.add_paragraph(texto, style="Heading 3")

    def h4(self, texto):
        return self.doc.add_paragraph(texto, style="Heading 4")

    def p(self, texto, sangria=True):
        estilo = "Cuerpo" if sangria else "CuerpoSinSangria"
        return self.doc.add_paragraph(texto, style=estilo)

    def ps(self, *textos):
        for t in textos:
            self.p(t)

    def negrita_p(self, etiqueta, texto):
        par = self.doc.add_paragraph(style="CuerpoSinSangria")
        r = par.add_run(etiqueta)
        r.bold = True
        par.add_run(" " + texto)
        return par

    def vinetas(self, items):
        for it in items:
            if isinstance(it, tuple):
                par = self.doc.add_paragraph(style="List Bullet")
                r = par.add_run(it[0])
                r.bold = True
                par.add_run(" " + it[1])
            else:
                self.doc.add_paragraph(it, style="List Bullet")

    def numerada(self, items):
        for it in items:
            self.doc.add_paragraph(it, style="List Number")

    def nota(self, texto):
        p = self.doc.add_paragraph(texto, style="Nota")
        _sombra_parrafo(p, "FBFCFF")
        return p

    def codigo(self, texto, titulo=None):
        if titulo:
            t = self.doc.add_paragraph(titulo, style="TablaTexto")
            t.runs[0].bold = True
            _mantener_con_siguiente(t)
        tabla = self.doc.add_table(rows=1, cols=1)
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordes_tabla(tabla, "DDE2EC", 4)
        _layout_fijo(tabla)
        tabla.columns[0].width = Cm(15.5)
        tabla.cell(0, 0).width = Cm(15.5)
        celda = tabla.cell(0, 0)
        _sombra(celda, "FAFBFE")
        celda.paragraphs[0].text = ""
        primera = True
        for linea in texto.split("\n"):
            par = celda.paragraphs[0] if primera else celda.add_paragraph()
            par.style = self.doc.styles["Codigo"]
            par.text = linea if linea.strip() else " "
            primera = False
        self.doc.add_paragraph(style="TablaTexto")
        return tabla

    def salto_pagina(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def espacio(self, n=1):
        for _ in range(n):
            self.doc.add_paragraph(style="CuerpoSinSangria")

    # ------------------------------------------------------------------ figuras
    def figura(self, archivo, descripcion, fuente="Elaboracion propia.", ancho=13.5):
        ruta = os.path.join(IMG, archivo)
        self.fig_n += 1
        num = f"{self.capitulo}.{self.fig_n}"

        cont = self.doc.add_paragraph()
        cont.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cont.paragraph_format.space_before = Pt(10)
        cont.paragraph_format.space_after = Pt(4)
        cont.paragraph_format.line_spacing = 1.0
        cont.paragraph_format.keep_with_next = True
        if os.path.exists(ruta):
            cont.add_run().add_picture(ruta, width=Cm(ancho))
        else:
            cont.add_run(f"[ Imagen no encontrada: {archivo} ]")

        cap = self.doc.add_paragraph(style="Figura")
        cap.add_run(f"Figura {num}")
        cap.add_run().add_break()
        cap.add_run(descripcion)
        cap.paragraph_format.keep_with_next = True

        f = self.doc.add_paragraph(f"Fuente: {fuente}", style="FiguraFuente")
        self.indice_figuras.append((num, descripcion))
        return num

    # ------------------------------------------------------------------ tablas
    def tabla(self, descripcion, encabezados, filas, fuente="Elaboracion propia.",
              anchos=None, fs=9.5, alineacion_col=None, ancho_total=15.5):
        self.tab_n += 1
        num = f"{self.capitulo}.{self.tab_n}"

        cap = self.doc.add_paragraph(style="Tabla")
        cap.add_run(f"Tabla {num}")
        cap.add_run().add_break()
        cap.add_run(descripcion)
        cap.paragraph_format.keep_with_next = True

        t = self.doc.add_table(rows=1, cols=len(encabezados))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _bordes_tabla(t)
        _layout_fijo(t)

        hdr = t.rows[0]
        _repetir_encabezado(hdr)
        for i, texto in enumerate(encabezados):
            c = hdr.cells[i]
            _sombra(c, NAVY)
            par = c.paragraphs[0]
            par.style = self.doc.styles["TablaEncabezado"]
            r = par.add_run(str(texto))
            r.font.size = Pt(fs)

        for k, fila in enumerate(filas):
            row = t.add_row()
            if k % 2 == 1:
                for c in row.cells:
                    _sombra(c, "F7F9FC")
            for i, valor in enumerate(fila):
                if i >= len(encabezados):
                    break
                c = row.cells[i]
                par = c.paragraphs[0]
                par.style = self.doc.styles["TablaTexto"]
                if alineacion_col and alineacion_col[i] == "c":
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif alineacion_col and alineacion_col[i] == "j":
                    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                partes = str(valor).split("\n")
                r = par.add_run(partes[0])
                r.font.size = Pt(fs)
                for extra in partes[1:]:
                    par.add_run().add_break()
                    rr = par.add_run(extra)
                    rr.font.size = Pt(fs)

        if not anchos:
            anchos = [1] * len(encabezados)
        if anchos:
            total = sum(anchos)
            for i, a in enumerate(anchos):
                ancho_cm = Cm(ancho_total * a / total)
                for row in t.rows:
                    row.cells[i].width = ancho_cm

        f = self.doc.add_paragraph(f"Fuente: {fuente}", style="FiguraFuente")
        self.indice_tablas.append((num, descripcion))
        return num

    # ------------------------------------------------------------------ caso de prueba
    def caso_prueba(self, cp):
        """Renderiza un caso de prueba completo como tabla de dos columnas."""
        self.tab_n += 1
        num = f"{self.capitulo}.{self.tab_n}"

        cap = self.doc.add_paragraph(style="Tabla")
        cap.add_run(f"Tabla {num}")
        cap.add_run().add_break()
        cap.add_run(f"Caso de prueba {cp['id']} — {cp['nombre']}.")
        cap.paragraph_format.keep_with_next = True

        pasos = "\n".join(f"{i+1}. {s}" for i, s in enumerate(cp["pasos"]))
        estado = cp["estado"]

        campos = [
            ("ID del caso", cp["id"]),
            ("Nombre", cp["nombre"]),
            ("Modulo", cp["modulo"]),
            ("Requerimiento", cp["rf"]),
            ("Tipo de prueba", cp["tipo"]),
            ("Tecnica de diseno", cp["tecnica"]),
            ("Prioridad", cp["prioridad"]),
            ("Objetivo", cp["objetivo"]),
            ("Precondiciones", cp["precondiciones"]),
            ("Datos de entrada", cp["datos"]),
            ("Pasos de ejecucion", pasos),
            ("Resultado esperado", cp["esperado"]),
            ("Resultado obtenido", cp["obtenido"]),
            ("Estado", estado),
            ("Defecto asociado", cp.get("defecto", "No aplica")),
            ("Ejecutado por", cp.get("ejecutor", "Carrete Ruiz Abril Carolina")),
            ("Fecha de ejecucion", cp.get("fecha", "17/11/2025")),
        ]

        t = self.doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _bordes_tabla(t)
        _layout_fijo(t)

        for etiqueta, valor in campos:
            row = t.add_row()
            _no_partir_fila(row)
            c0, c1 = row.cells
            _sombra(c0, NAVY_SOFT)
            p0 = c0.paragraphs[0]
            p0.style = self.doc.styles["TablaTexto"]
            r0 = p0.add_run(etiqueta)
            r0.bold = True
            r0.font.size = Pt(9.5)

            p1 = c1.paragraphs[0]
            p1.style = self.doc.styles["TablaTexto"]
            lineas = str(valor).split("\n")
            r1 = p1.add_run(lineas[0])
            r1.font.size = Pt(9.5)
            if etiqueta == "Estado":
                r1.bold = True
                r1.font.color.rgb = RGBColor.from_string(
                    "1E7E34" if estado.startswith("Aprobado") else
                    ("8B1E1E" if estado.startswith("Fallido") else "9A6B00"))
            for extra in lineas[1:]:
                p1.add_run().add_break()
                rr = p1.add_run(extra)
                rr.font.size = Pt(9.5)

            c0.width = Cm(4.2)
            c1.width = Cm(11.3)

        self.doc.add_paragraph(f"Fuente: Elaboracion propia.", style="FiguraFuente")
        self.indice_tablas.append((num, f"Caso de prueba {cp['id']} — {cp['nombre']}."))
        return num

    # ------------------------------------------------------------------ indices
    def indice(self, titulo, instruccion):
        p = self.doc.add_paragraph(titulo, style="TituloSuelto")
        _borde_inferior(p, NAVY, 8)
        par = self.doc.add_paragraph()
        par.paragraph_format.line_spacing = 1.5
        _campo(par, instruccion,
               "Coloque el cursor aqui, presione clic derecho y elija «Actualizar campos» "
               "para generar este indice.")

    def toc_general(self):
        self.indice("Indice", 'TOC \\o "1-3" \\h \\z \\u')

    def toc_figuras(self):
        self.indice("Indice de figuras", 'TOC \\h \\z \\t "Figura,1"')

    def toc_tablas(self):
        self.indice("Indice de tablas", 'TOC \\h \\z \\t "Tabla,1"')

    # ------------------------------------------------------------------ guardado
    def guardar(self, ruta):
        # Pide a Word que actualice los campos (indices) al abrir el archivo
        settings = self.doc.settings.element
        uf = settings.find(qn("w:updateFields"))
        if uf is None:
            settings.append(_el("w:updateFields", val="true"))
        self.doc.save(ruta)
        return ruta
